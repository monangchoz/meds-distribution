# isolated_dimension_export.py
import pandas as pd
import numpy as np
import os
import glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def extract_cluster_info(filename):
    """
    Extract cluster information from filename
    Returns: number of clusters (1, 3, or 5) or None if not found
    """
    filename_upper = filename.upper()
    if 'NCL_1' in filename_upper:
        return 1
    elif 'NCL_3' in filename_upper:
        return 3
    elif 'NCL_5' in filename_upper:
        return 5
    return None

def extract_data_type(filename):
    """
    Extract data type (historical or generated) from filename
    Returns: 'historical', 'generated', or None if not found
    """
    filename_upper = filename.upper()
    if 'HISTORICAL' in filename_upper:
        return 'historical'
    elif 'GENERATED' in filename_upper:
        return 'generated'
    return None

def load_and_process_csv_files(base_path):
    """
    Load all CSV files from algorithm folders and organize by multiple dimensions
    """
    # Multi-dimensional dictionary structure
    data_by_location = {
        'JK2': {},
        'MKS': {},
        'SBY': {}
    }
    
    data_by_cluster = {1: {}, 3: {}, 5: {}}
    data_by_datatype = {'historical': {}, 'generated': {}}
    
    # Define algorithm folders and their display names
    algorithm_folders = {
        'avns': 'AVNS',
        'ga': 'GA', 
        'de': 'DE',
        'brkga': 'BRKGA',
        'pso': 'PSO'
    }
    
    results_path = os.path.join(base_path, 'results')
    
    if not os.path.exists(results_path):
        print(f"Results folder not found at: {results_path}")
        return data_by_location, data_by_cluster, data_by_datatype
    
    # Process each algorithm folder
    for folder_name, algorithm_name in algorithm_folders.items():
        algorithm_path = os.path.join(results_path, folder_name)
        
        if not os.path.exists(algorithm_path):
            print(f"Algorithm folder not found: {algorithm_path}")
            continue
            
        # Get all CSV files in this algorithm folder
        csv_files = glob.glob(os.path.join(algorithm_path, "*.csv"))
        
        for file_path in csv_files:
            filename = os.path.basename(file_path)
            filename_upper = filename.upper()
            
            # Extract all dimensions
            location = None
            if 'JK2' in filename_upper:
                location = 'JK2'
            elif 'MKS' in filename_upper:
                location = 'MKS'
            elif 'SBY' in filename_upper:
                location = 'SBY'
            
            cluster_count = extract_cluster_info(filename)
            data_type = extract_data_type(filename)
            
            try:
                # Read the CSV file
                df = pd.read_csv(file_path, header=None, names=['total_cost', 'running_time'])
                
                # Store in location-based structure
                if location:
                    if algorithm_name not in data_by_location[location]:
                        data_by_location[location][algorithm_name] = {'total_cost': [], 'running_time': []}
                    data_by_location[location][algorithm_name]['total_cost'].extend(df['total_cost'].tolist())
                    data_by_location[location][algorithm_name]['running_time'].extend(df['running_time'].tolist())
                
                # Store in cluster-based structure
                if cluster_count is not None:
                    if algorithm_name not in data_by_cluster[cluster_count]:
                        data_by_cluster[cluster_count][algorithm_name] = {'total_cost': [], 'running_time': []}
                    data_by_cluster[cluster_count][algorithm_name]['total_cost'].extend(df['total_cost'].tolist())
                    data_by_cluster[cluster_count][algorithm_name]['running_time'].extend(df['running_time'].tolist())
                
                # Store in data type-based structure
                if data_type:
                    if algorithm_name not in data_by_datatype[data_type]:
                        data_by_datatype[data_type][algorithm_name] = {'total_cost': [], 'running_time': []}
                    data_by_datatype[data_type][algorithm_name]['total_cost'].extend(df['total_cost'].tolist())
                    data_by_datatype[data_type][algorithm_name]['running_time'].extend(df['running_time'].tolist())
                
                # Enhanced logging
                cluster_info = f", Clusters: {cluster_count}" if cluster_count else ""
                datatype_info = f", Data: {data_type}" if data_type else ""
                print(f"Processed {filename} -> Location: {location}, Algorithm: {algorithm_name}{cluster_info}{datatype_info}, Rows: {len(df)}")
                
            except Exception as e:
                print(f"Error processing {filename}: {e}")
    
    return data_by_location, data_by_cluster, data_by_datatype

def apply_excel_styling(sheet, headers, title):
    """
    Apply consistent styling to Excel sheets
    """
    # Setup styling
    header_font = Font(bold=True, size=12)
    title_font = Font(bold=True, size=14)
    header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                        top=Side(style='thin'), bottom=Side(style='thin'))
    center_align = Alignment(horizontal='center', vertical='center')
    
    # Add title
    sheet['A1'] = title
    sheet['A1'].font = title_font
    sheet.merge_cells(f'A1:{get_column_letter(len(headers))}1')
    sheet['A1'].alignment = center_align
    
    # Create headers
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = center_align
    
    return sheet

def create_isolated_analysis_excel(data_by_location, data_by_cluster, data_by_datatype, output_filename="isolated_algorithm_analysis.xlsx"):
    """
    Create single Excel file with 3 sheets for isolated dimension analysis
    """
    try:
        wb = Workbook()
        
        # === SHEET 1: DATA TYPE ANALYSIS ===
        sheet1 = wb.active
        sheet1.title = "Data Type Analysis"
        
        headers1 = ["Data Type", "Algorithm", "Mean Total Cost", "Std Dev Total Cost", 
                   "Mean Running Time (s)", "Std Dev Running Time (s)"]
        
        apply_excel_styling(sheet1, headers1, "Algorithm Performance by Data Type")
        
        row = 4
        for data_type in ['historical', 'generated']:
            if data_type in data_by_datatype:
                for algorithm in sorted(data_by_datatype[data_type].keys()):
                    cost_data = data_by_datatype[data_type][algorithm]['total_cost']
                    time_data = data_by_datatype[data_type][algorithm]['running_time']
                    
                    if cost_data and time_data:
                        mean_cost = np.mean(cost_data)
                        std_cost = np.std(cost_data)
                        mean_time = np.mean(time_data)
                        std_time = np.std(time_data)
                        
                        # Add data to sheet
                        data_row = [data_type.capitalize(), algorithm, mean_cost, std_cost, mean_time, std_time]
                        for col, value in enumerate(data_row, 1):
                            cell = sheet1.cell(row=row, column=col, value=value)
                            cell.border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                               top=Side(style='thin'), bottom=Side(style='thin'))
                            
                            # Format numbers
                            if col >= 3:
                                if col in [3, 4]:  # Cost columns
                                    cell.number_format = '#,##0.00'
                                else:  # Time columns
                                    cell.number_format = '0.0000'
                        row += 1
                
                # Add separator between data types
                row += 1
        
        # Adjust column widths for sheet1
        column_widths1 = [15, 12, 18, 20, 20, 22]
        for col, width in enumerate(column_widths1, 1):
            sheet1.column_dimensions[get_column_letter(col)].width = width
        
        # === SHEET 2: LOCATION ANALYSIS ===
        sheet2 = wb.create_sheet("Location Analysis")
        
        headers2 = ["Location", "Algorithm", "Mean Total Cost", "Std Dev Total Cost", 
                   "Mean Running Time (s)", "Std Dev Running Time (s)"]
        
        apply_excel_styling(sheet2, headers2, "Algorithm Performance by Location")
        
        location_names = {'JK2': 'Jakarta 2', 'MKS': 'Makassar', 'SBY': 'Surabaya'}
        
        row = 4
        for location in ['JK2', 'MKS', 'SBY']:
            if location in data_by_location and data_by_location[location]:
                for algorithm in sorted(data_by_location[location].keys()):
                    cost_data = data_by_location[location][algorithm]['total_cost']
                    time_data = data_by_location[location][algorithm]['running_time']
                    
                    if cost_data and time_data:
                        mean_cost = np.mean(cost_data)
                        std_cost = np.std(cost_data)
                        mean_time = np.mean(time_data)
                        std_time = np.std(time_data)
                        
                        # Add data to sheet
                        data_row = [location_names[location], algorithm, mean_cost, std_cost, mean_time, std_time]
                        for col, value in enumerate(data_row, 1):
                            cell = sheet2.cell(row=row, column=col, value=value)
                            cell.border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                               top=Side(style='thin'), bottom=Side(style='thin'))
                            
                            # Format numbers
                            if col >= 3:
                                if col in [3, 4]:  # Cost columns
                                    cell.number_format = '#,##0.00'
                                else:  # Time columns
                                    cell.number_format = '0.0000'
                        row += 1
                
                # Add separator between locations
                row += 1
        
        # Adjust column widths for sheet2
        column_widths2 = [15, 12, 18, 20, 20, 22]
        for col, width in enumerate(column_widths2, 1):
            sheet2.column_dimensions[get_column_letter(col)].width = width
        
        # === SHEET 3: CLUSTER ANALYSIS ===
        sheet3 = wb.create_sheet("Cluster Analysis")
        
        headers3 = ["Cluster Count", "Algorithm", "Mean Total Cost", "Std Dev Total Cost", 
                   "Mean Running Time (s)", "Std Dev Running Time (s)"]
        
        apply_excel_styling(sheet3, headers3, "Algorithm Performance by Cluster Count")
        
        row = 4
        for cluster_count in [1, 3, 5]:
            if cluster_count in data_by_cluster and data_by_cluster[cluster_count]:
                for algorithm in sorted(data_by_cluster[cluster_count].keys()):
                    cost_data = data_by_cluster[cluster_count][algorithm]['total_cost']
                    time_data = data_by_cluster[cluster_count][algorithm]['running_time']
                    
                    if cost_data and time_data:
                        mean_cost = np.mean(cost_data)
                        std_cost = np.std(cost_data)
                        mean_time = np.mean(time_data)
                        std_time = np.std(time_data)
                        
                        # Add data to sheet
                        data_row = [cluster_count, algorithm, mean_cost, std_cost, mean_time, std_time]
                        for col, value in enumerate(data_row, 1):
                            cell = sheet3.cell(row=row, column=col, value=value)
                            cell.border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                               top=Side(style='thin'), bottom=Side(style='thin'))
                            
                            # Format numbers
                            if col >= 3:
                                if col in [3, 4]:  # Cost columns
                                    cell.number_format = '#,##0.00'
                                else:  # Time columns
                                    cell.number_format = '0.0000'
                        row += 1
                
                # Add separator between cluster counts
                row += 1
        
        # Adjust column widths for sheet3
        column_widths3 = [15, 12, 18, 20, 20, 22]
        for col, width in enumerate(column_widths3, 1):
            sheet3.column_dimensions[get_column_letter(col)].width = width
        
        # Save the workbook
        wb.save(output_filename)
        print(f"✅ Excel file with isolated analysis created: {output_filename}")
        return True
        
    except Exception as e:
        print(f"❌ Error creating Excel file: {e}")
        return False

def main():
    """
    Main function to create Excel file with 3 isolated analysis sheets
    """
    print("Isolated Dimension Analysis - Single Excel File")
    print("=" * 50)
    print("This script will create 1 Excel file with 3 sheets:")
    print("1. Data Type Analysis    - Historical vs Generated")
    print("2. Location Analysis     - JK2, MKS, SBY")
    print("3. Cluster Analysis      - 1, 3, 5 clusters")
    print("=" * 50)
    
    # Get base path
    base_path = os.getcwd()
    print(f"Looking for CSV files in: {base_path}/results/")
    print("-" * 50)
    
    # Load and process data
    data_by_location, data_by_cluster, data_by_datatype = load_and_process_csv_files(base_path)
    
    # Check if data was loaded
    locations_with_data = [loc for loc in ['JK2', 'MKS', 'SBY'] if data_by_location[loc]]
    clusters_with_data = [cluster for cluster in [1, 3, 5] if data_by_cluster[cluster]]
    datatypes_with_data = [dt for dt in ['historical', 'generated'] if data_by_datatype[dt]]
    
    print(f"\nData found for:")
    print(f"  - Locations: {locations_with_data}")
    print(f"  - Clusters: {clusters_with_data}")
    print(f"  - Data Types: {datatypes_with_data}")
    print("-" * 50)
    
    if not any([locations_with_data, clusters_with_data, datatypes_with_data]):
        print("❌ No data found! Please check the folder structure and file naming conventions.")
        return
    
    # Get output filename
    default_filename = "isolated_algorithm_analysis.xlsx"
    print(f"\nEnter filename for Excel output:")
    print(f"(Press Enter to use default: {default_filename})")
    excel_filename = input("> ").strip()
    
    if not excel_filename:
        excel_filename = default_filename
    if not excel_filename.endswith('.xlsx'):
        excel_filename += '.xlsx'
    
    # Create Excel file with isolated analysis
    if create_isolated_analysis_excel(data_by_location, data_by_cluster, data_by_datatype, excel_filename):
        print(f"\n🎉 Successfully created Excel file: {excel_filename}")
        print("\nFile contains 3 sheets:")
        print("  📊 Data Type Analysis - Compare Historical vs Generated data")
        print("  📍 Location Analysis - Compare JK2 vs MKS vs SBY performance") 
        print("  🎯 Cluster Analysis - Compare 1 vs 3 vs 5 cluster performance")
        print("\nEach sheet shows:")
        print("  - Mean and Standard Deviation for Total Cost")
        print("  - Mean and Standard Deviation for Running Time")
        print("  - Professional formatting with borders and styling")

if __name__ == "__main__":
    main()